import csv
import io
import re
from datetime import date, datetime

from django.core.exceptions import PermissionDenied, ValidationError
from django.db import models, transaction

from apps.auditlog.models import AuditEvent
from apps.tenancy.models import Membership

from .models import Asset, CriticalService, Site


FIELD_ALIASES = {
    "asset_tag": {"assettag", "assetid", "assetcode", "inventoryid", "inventorycode", "tag", "deviceid"},
    "name": {"name", "assetname", "devicename", "hostname", "systemname"},
    "kind": {"kind", "type", "assettype", "devicetype", "category", "assetcategory"},
    "owner_username": {"owner", "ownerusername", "assetowner", "businessowner", "responsibleperson", "responsibleowner"},
    "custodian_username": {"custodian", "custodianusername", "assetcustodian", "technicalowner", "technicalcustodian"},
    "criticality": {"criticality", "assetcriticality", "businesscriticality", "impact", "priority"},
    "site_code": {"site", "sitecode", "siteid", "location", "locationcode", "pop", "office"},
    "service_code": {"service", "servicecode", "businessservice", "criticalservice", "serviceid"},
    "is_cii": {"iscii", "cii", "ciistatus", "criticalinfrastructure", "criticalasset"},
    "classification": {"classification", "dataclassification", "securityclassification", "assetclassification"},
    "lifecycle": {"lifecycle", "lifecyclestatus", "assetstatus", "status"},
    "manufacturer": {"manufacturer", "vendor", "make", "brand"},
    "model": {"model", "modelnumber", "productmodel"},
    "serial_number": {"serial", "serialnumber", "serialno", "servicetag"},
    "source_reference": {"sourcereference", "source", "inventorysource", "cmdbreference", "reference"},
    "review_due_at": {"reviewdue", "reviewdate", "nextreview", "reviewdueat"},
}
REQUIRED_FIELDS = {"asset_tag", "name", "kind", "owner_username", "custodian_username", "criticality"}


def _key(value):
    return re.sub(r"[^a-z0-9]", "", str(value or "").strip().lower())


def _field_map(headers):
    mapped = {}
    for header in headers:
        normalized = _key(header)
        for field, aliases in FIELD_ALIASES.items():
            if normalized in aliases and field not in mapped:
                mapped[field] = header
                break
    missing = REQUIRED_FIELDS - mapped.keys()
    if missing:
        raise ValidationError("Required details could not be identified: " + ", ".join(sorted(missing)) + ". Rename only these unclear columns or use a recognized heading such as Asset ID, Asset Name, Type, Owner, Custodian and Criticality.")
    return mapped


def _sheet_rows(*, filename, content):
    if not content:
        raise ValidationError("The inventory file is empty.")
    extension = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    # Trust the file signature over a misleading extension. Excel exports are
    # commonly renamed to .csv without being converted to text.
    if content.startswith(b"PK\x03\x04"):
        extension = "xlsx"
    elif content.startswith(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"):
        extension = "xls"
    if extension == "csv":
        text = None
        for encoding in ("utf-8-sig", "utf-32", "utf-16", "utf-16-le", "utf-16-be", "cp1252", "latin-1"):
            try:
                text = content.decode(encoding)
                if "\x00" in text:
                    text = None
                    continue
                break
            except UnicodeDecodeError:
                continue
        if text is None:
            raise ValidationError("The CSV text encoding could not be detected.")
        try:
            dialect = csv.Sniffer().sniff(text[:8192], delimiters=",;\t|")
        except csv.Error:
            dialect = csv.excel
        reader = csv.DictReader(io.StringIO(text), dialect=dialect)
        return reader.fieldnames or [], list(reader)
    if extension == "xlsx":
        from openpyxl import load_workbook
        try:
            sheet = load_workbook(io.BytesIO(content), read_only=True, data_only=True).active
            values = list(sheet.iter_rows(values_only=True))
        except Exception as exc:
            raise ValidationError("XLSX file could not be read.") from exc
    elif extension == "xls":
        import xlrd
        try:
            sheet = xlrd.open_workbook(file_contents=content).sheet_by_index(0)
            values = [sheet.row_values(index) for index in range(sheet.nrows)]
        except Exception as exc:
            raise ValidationError("XLS file could not be read.") from exc
    else:
        raise ValidationError("Upload a CSV, XLSX or XLS asset inventory file.")
    if not values:
        raise ValidationError("The spreadsheet is empty.")
    headers = [str(value or "").strip() for value in values[0]]
    rows = [dict(zip(headers, row)) for row in values[1:] if any(value not in (None, "") for value in row)]
    return headers, rows


def _enum(value, choices, aliases=None):
    normalized = _key(value)
    aliases = aliases or {}
    if normalized in aliases:
        return aliases[normalized]
    for code, label in choices:
        if normalized in {_key(code), _key(label)}:
            return code
    return str(value or "").strip().upper()


def _member_user(*, tenant, value, label):
    wanted = str(value or "").strip().casefold()
    matches = []
    for membership in tenant.memberships.filter(is_active=True).select_related("user"):
        user = membership.user
        candidates = {user.username.casefold(), user.email.casefold(), user.get_full_name().strip().casefold()}
        if wanted and wanted in candidates:
            matches.append(user)
    if len(matches) != 1:
        raise ValidationError(f"{label} '{value}' must match exactly one active tenant username, email or full name.")
    return matches[0]


def require_inventory_operator(*, tenant, actor):
    membership = Membership.objects.filter(tenant=tenant, user=actor, is_active=True).first()
    if not membership or membership.role not in {Membership.Role.OWNER, Membership.Role.COMPLIANCE_LEAD}:
        raise PermissionDenied("Tenant owner or compliance lead access is required.")


def save_inventory_object(*, obj, actor):
    require_inventory_operator(tenant=obj.tenant, actor=actor)
    obj.full_clean()
    obj.save()
    AuditEvent.objects.create(tenant=obj.tenant, actor=actor, action=f"inventory.{obj._meta.model_name}_created", object_type=obj.__class__.__name__, object_id=str(obj.pk), metadata={})
    return obj


@transaction.atomic
def import_assets_file(*, tenant, actor, filename, content):
    require_inventory_operator(tenant=tenant, actor=actor)
    headers, source_rows = _sheet_rows(filename=filename, content=content)
    mapping = _field_map(headers)
    rows, errors, seen_tags = [], [], set()
    for number, source_row in enumerate(source_rows, start=2):
        try:
            row = {field: source_row.get(header, "") for field, header in mapping.items()}
            row = {field: "" if value is None else str(value).strip() for field, value in row.items()}
            owner = _member_user(tenant=tenant, value=row["owner_username"], label="Owner")
            custodian = _member_user(tenant=tenant, value=row["custodian_username"], label="Custodian")
            site_code, service_code = row.get("site_code", "").strip(), row.get("service_code", "").strip()
            site = Site.objects.filter(tenant=tenant).filter(models.Q(code__iexact=site_code) | models.Q(name__iexact=site_code)).first() if site_code else None
            service = CriticalService.objects.filter(tenant=tenant).filter(models.Q(code__iexact=service_code) | models.Q(name__iexact=service_code)).first() if service_code else None
            if site_code and not site:
                raise ValidationError(f"Unknown site_code {site_code}.")
            if service_code and not service:
                raise ValidationError(f"Unknown service_code {service_code}.")
            tag = row["asset_tag"].strip()
            if tag in seen_tags:
                raise ValidationError(f"Duplicate asset_tag {tag} in this batch.")
            seen_tags.add(tag)
            kind = _enum(row["kind"], Asset.Kind.choices, {"router": Asset.Kind.NETWORK_DEVICE, "switch": Asset.Kind.NETWORK_DEVICE, "olt": Asset.Kind.NETWORK_DEVICE, "firewall": Asset.Kind.SECURITY_DEVICE, "vm": Asset.Kind.SERVER})
            criticality = _enum(row["criticality"], CriticalService.Criticality.choices)
            classification = _enum(row.get("classification", "INTERNAL"), Asset.Classification.choices)
            lifecycle = _enum(row.get("lifecycle", "ACTIVE"), Asset.Lifecycle.choices)
            review_value = source_row.get(mapping.get("review_due_at", ""), "")
            review_due = review_value.date() if isinstance(review_value, datetime) else review_value if isinstance(review_value, date) else (datetime.fromisoformat(str(review_value)).date() if review_value else None)
            obj = Asset(tenant=tenant, asset_tag=tag, name=row["name"], kind=kind, owner=owner, custodian=custodian, criticality=criticality, site=site, service=service, is_cii=row.get("is_cii", "").lower() in {"1", "true", "yes", "y", "cii"}, classification=classification, lifecycle=lifecycle, manufacturer=row.get("manufacturer", ""), model=row.get("model", ""), serial_number=row.get("serial_number", ""), source_reference=row.get("source_reference", ""), review_due_at=review_due, created_by=actor)
            obj.full_clean()
            rows.append(obj)
        except Exception as exc:
            errors.append(f"Row {number}: {exc}")
    if errors:
        raise ValidationError(errors)
    Asset.objects.bulk_create(rows)
    AuditEvent.objects.create(tenant=tenant, actor=actor, action="inventory.assets_imported", object_type="Asset", object_id="batch", metadata={"count": len(rows)})
    return len(rows)


def import_assets_csv(*, tenant, actor, content):
    """Backward-compatible wrapper for existing callers."""
    return import_assets_file(tenant=tenant, actor=actor, filename="assets.csv", content=content)
